"""Groq-powered resume refinement and DOCX export."""

import io
import json
import re

from docx import Document

from src.services.groq_service import ask_groq


def refine_resume(resume_text: str, profile: dict) -> dict:
    """Improve resume readability without inventing candidate information."""

    if not isinstance(resume_text, str) or not resume_text.strip():
        raise ValueError(
            "Resume text is empty. Analyze a resume before refining it."
        )

    prompt = f"""
You are a careful professional resume editor.

Improve the candidate's resume for clarity, grammar, structure,
professional presentation, and ATS readability.

STRICT RULES:
1. Never invent employers, dates, degrees, certifications, skills,
   metrics, achievements, or responsibilities.
2. Preserve the original factual meaning.
3. Do not claim guaranteed ATS success, interviews, or employment.
4. If important information is missing, list it under missing_information.
5. Use standard resume headings and concise bullet points.
6. Do not add facts merely because they would make the resume stronger.
7. Return only valid JSON. Do not include Markdown code fences.

Return a JSON object with exactly these keys:
{{
  "refined_resume_markdown": "The complete refined resume",
  "changes": ["Description of meaningful changes"],
  "missing_information": ["Information the user should verify or add"]
}}

Extracted profile:
{json.dumps(profile, ensure_ascii=False)}

Original resume:
<resume>
{resume_text[:14000]}
</resume>
"""

    response = ask_groq(prompt, json_mode=True)

    try:
        result = json.loads(response)
    except json.JSONDecodeError as exc:
        raise ValueError(
            "Groq returned invalid JSON. Please try again."
        ) from exc

    required_keys = {
        "refined_resume_markdown",
        "changes",
        "missing_information",
    }

    if not isinstance(result, dict):
        raise ValueError("Unexpected response from Groq.")

    if not required_keys.issubset(result.keys()):
        raise ValueError(
            "Groq response is missing required resume fields."
        )

    refined_text = result["refined_resume_markdown"]

    if not isinstance(refined_text, str) or not refined_text.strip():
        raise ValueError("Groq returned an empty refined resume.")

    for key in ("changes", "missing_information"):
        if not isinstance(result[key], list):
            result[key] = []

    return result


def create_docx(markdown_text: str) -> bytes:
    """Convert simple Markdown headings and bullets into a DOCX file."""

    if not isinstance(markdown_text, str) or not markdown_text.strip():
        raise ValueError("There is no resume content to export.")

    document = Document()

    for line in markdown_text.splitlines():
        stripped = line.strip()

        if not stripped:
            continue

        if stripped.startswith("### "):
            document.add_heading(stripped[4:], level=3)

        elif stripped.startswith("## "):
            document.add_heading(stripped[3:], level=2)

        elif stripped.startswith("# "):
            document.add_heading(stripped[2:], level=1)

        elif re.match(r"^[-*•]\s+", stripped):
            bullet_text = re.sub(r"^[-*•]\s+", "", stripped)
            document.add_paragraph(
                bullet_text,
                style="List Bullet",
            )

        else:
            document.add_paragraph(stripped)

    buffer = io.BytesIO()
    document.save(buffer)

    return buffer.getvalue()